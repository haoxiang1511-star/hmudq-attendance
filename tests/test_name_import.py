# -*- coding: utf-8 -*-
import os, sys, tempfile, unittest
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from name_import import parse_names, merge_names, clean_cell, parse_names_detail, normalize_name, disambiguate

N = ["赵甲", "赵乙", "赵丙", "赵丁", "赵戊", "李甲", "李戊", "迪丽热巴·迪力木拉提", "欧阳娜娜", "周戊"]

def big(n):
    fam = "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜"
    giv = "伟芳娜敏静丽强磊军洋勇艳杰娟涛明超秀霞平刚桂英华玉兰萍红"
    out = []
    for i in range(n):
        out.append(fam[i % len(fam)] + giv[(i // len(fam)) % len(giv)] + giv[(i * 7) % len(giv)])
    return list(dict.fromkeys(out))

class T(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()
    def save(self, wb, name="t.xlsx"):
        p = os.path.join(self.d, name); wb.save(p); return p

    def test_group_list(self):
        wb = openpyxl.Workbook(); ws = wb.active
        ws["A1"] = "24示例3班 课程分组名单"; ws.merge_cells("A1:C1")
        for i, n in enumerate(N[:8]):
            ws.cell(row=i + 2, column=1, value=f"第{i // 4 + 1}组" if i % 4 == 0 else None)
            ws.cell(row=i + 2, column=2, value=f"{n}（组长）" if i % 4 == 0 else n)
        ws.merge_cells("A2:A5"); ws.merge_cells("A6:A9")
        names, info, err = parse_names(self.save(wb))
        self.assertIsNone(err); self.assertEqual(names, N[:8])

    def test_header_with_extra_cols(self):
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["2026年寒假离校确认表"]); ws.append(["序号", "学号", "姓名", "性别", "电话", "宿舍", "去向", "本人签字"])
        for i, n in enumerate(N):
            ws.append([i + 1, f"2024{i:04d}", n, "男", "13800000000", "3-201", "回家", n])
        names, info, err = parse_names(self.save(wb))
        self.assertEqual(names, N); self.assertIn("表头", info)

    def test_no_header_density(self):
        wb = openpyxl.Workbook(); ws = wb.active
        for i, n in enumerate(N):
            ws.append([i + 1, "是" if i % 2 else "否", n, "备注"])
        names, _, _ = parse_names(self.save(wb))
        self.assertEqual(names, N)

    def test_multi_sheet_and_prefix(self):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "说明"
        ws["A1"] = "活动报名说明"; ws["A2"] = "请按时参加"
        ws2 = wb.create_sheet("报名")
        ws2.append(["活动报名表"]); ws2.append(["名字", "联系方式"])
        ws2.append(["组长：赵甲", "1"]); ws2.append(["1. 赵乙", "2"]); ws2.append(["赵丙 ", "3"])
        ws2.append(["赵丙", "dup"]); ws2.append(["合计", "3人"])
        names, info, _ = parse_names(self.save(wb))
        self.assertEqual(names, ["赵甲", "赵乙", "赵丙"]); self.assertIn("报名", info)

    def test_large_class_150(self):
        b = big(150)
        wb = openpyxl.Workbook(); ws = wb.active; ws.append(["学生姓名"])
        for n in b: ws.append([n])
        names, _, _ = parse_names(self.save(wb)); self.assertEqual(names, b)

    def test_csv_gbk(self):
        p = os.path.join(self.d, "a.csv")
        with open(p, "w", encoding="gb18030") as f:
            f.write("序号,姓名\n" + "\n".join(f"{i},{n}" for i, n in enumerate(N[:5])))
        names, _, err = parse_names(p); self.assertEqual(names, N[:5])

    def test_xls_message(self):
        p = os.path.join(self.d, "a.xls"); open(p, "wb").write(b"x")
        names, _, err = parse_names(p)
        self.assertTrue(err and ("xlsx" in err or "xls" in err))

    def test_empty(self):
        wb = openpyxl.Workbook(); wb.active.append(["序号", "日期"]); wb.active.append([1, "2026-10-01"])
        names, _, err = parse_names(self.save(wb)); self.assertEqual(names, []); self.assertIsNone(err)

    def test_clean_and_merge(self):
        for bad in ["序号", "男", "13800138000", "24示例3班", "第一组", "签字", "", None, 12]:
            self.assertEqual(clean_cell(bad), "")
        self.assertEqual(merge_names(["甲乙", "丙丁"], ["丙丁", "戊己"]), ["甲乙", "丙丁", "戊己"])

    # ---- v1.4.0 ----
    def test_normalize_name(self):
        self.assertEqual(normalize_name(" 张 三 "), "张三")
        self.assertEqual(normalize_name("张\u3000三"), "张三")
        self.assertEqual(normalize_name("  Tom   Smith "), "Tom Smith")
        self.assertEqual(normalize_name(None), "")

    def test_import_spaces_inside_chinese_name(self):
        wb = openpyxl.Workbook(); ws = wb.active
        ws["A1"] = "姓名"
        for i, n in enumerate(["赵 甲", "赵乙", "赵　丙", "赵丁"]):
            ws.cell(row=i + 2, column=1, value=n)
        names, info, err = parse_names(self.save(wb))
        self.assertIsNone(err)
        self.assertEqual(names, ["赵甲", "赵乙", "赵丙", "赵丁"])

    def test_import_reports_duplicates(self):
        wb = openpyxl.Workbook(); ws = wb.active
        ws["A1"] = "姓名"
        for i, n in enumerate(["张伟", "李娜", "张伟", "王芳", "张 伟"]):
            ws.cell(row=i + 2, column=1, value=n)
        names, info, err, dups = parse_names_detail(self.save(wb))
        self.assertEqual(names, ["张伟", "李娜", "王芳"])
        self.assertEqual(dups, {"张伟": 3})

    def test_disambiguate(self):
        self.assertEqual(disambiguate("张伟", ["张伟"]), "张伟(2)")
        self.assertEqual(disambiguate("张伟", ["张伟", "张伟(2)", "张伟(3)"]), "张伟(4)")

    def test_v140_name_rules(self):
        for ok in ["班婕", "组明", "阿依古丽买买提", "古丽娜尔·阿不都热合曼", "Tom Smith", "Mary-Jane O'Neil"]:
            self.assertEqual(clean_cell(ok), ok, ok)
        for bad in ["三班", "临床班", "签到表", "第一组", "课程分组名单", "请按时参加", "Name", "No", "ID"]:
            self.assertEqual(clean_cell(bad), "", bad)
        self.assertEqual(normalize_name("Tom   Smith"), "Tom Smith")

if __name__ == "__main__":
    unittest.main()
