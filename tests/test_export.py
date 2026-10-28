# -*- coding: utf-8 -*-
"""月度导出往返测试：随机生成考勤记录 -> 导出 -> 用下游 parse_att 风格解析器读回 -> 逐格比对。
v1.4.0 规则：没有记录的格子一律是「空课」（与点名页一致），不再自动算出勤。"""
import os, sys, random, tempfile, unittest, datetime
import openpyxl
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import main as app
_PARSE_ATT = os.environ.get("HMUDQ_PARSE_ATT", "")  # 可选：下游解析器路径，未设置则跳过对应比对
try:
    import importlib.util
    spec = importlib.util.spec_from_file_location("parse_att", _PARSE_ATT)
    _pa = importlib.util.module_from_spec(spec); sys.argv = sys.argv[:1]; spec.loader.exec_module(_pa)
    DOWNSTREAM = _pa.parse
except Exception:
    DOWNSTREAM = None

STATUSES = ["出勤", "旷课", "病假", "事假", "空课"]
FAM = "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜戚谢邹喻柏水窦章云苏潘葛"
GIV = "伟芳娜敏静丽强磊军洋勇艳杰娟涛明超秀霞平刚桂英华玉兰萍红"

def names(n, seed=0):
    out = []
    i = 0
    while len(out) < n:
        s = FAM[i % len(FAM)] + GIV[(i // len(FAM)) % len(GIV)] + GIV[(i * 7 + seed) % len(GIV)]
        if s not in out: out.append(s)
        i += 1
    return out

def own_parse(path):
    """返回 {(date, period_idx, name): value} 以及每个块的人名列表。"""
    ws = openpyxl.load_workbook(path).active
    cells, blocks = {}, []
    for r in range(1, ws.max_row + 1):
        for c0 in (1, 8, 15):
            if ws.cell(r, c0).value == "姓名":
                d = ws.cell(r, c0 + 1).value
                rr, ns = r + 1, []
                while rr <= ws.max_row and ws.cell(rr, c0).value not in (None, "姓名"):
                    n = ws.cell(rr, c0).value; ns.append(n)
                    for p in range(5):
                        key = (d, p, n)
                        assert key not in cells, f"重复格 {key}"
                        cells[key] = ws.cell(rr, c0 + 2 + p).value
                    rr += 1
                blocks.append((d, ns, [ws.cell(r, c0 + 2 + p).value for p in range(5)]))
    return cells, blocks

class ExportRoundTrip(unittest.TestCase):
    def fresh_db(self):
        d = tempfile.mkdtemp()
        app.DB_PATH = os.path.join(d, "t.db"); app._db_conn = None
        app.init_db()
        c = app.get_db_connection(); c.execute("DELETE FROM students"); c.commit()
        return c

    def run_scenario(self, n_initial, n_added, join_day, ym="2026-09", seed=1, days=31, reorder=False, delete_readd=False):
        rnd = random.Random(seed)
        import calendar
        days = min(days, calendar.monthrange(int(ym[:4]), int(ym[5:]))[1])
        conn = self.fresh_db()
        base = names(n_initial + n_added, seed)
        app.update_db_roster(base[:n_initial])
        join = f"{ym}-{join_day:02d}"
        conn.execute("UPDATE students SET joined_date=NULL"); conn.commit()
        app.update_db_roster(base)  # 追加新生
        conn.execute("UPDATE students SET joined_date=? WHERE name IN (%s)" % ",".join("?" * n_added), [join] + base[n_initial:]); conn.commit()
        if reorder:
            r = base[:]; rnd.shuffle(r); app.update_db_roster(r); base = r
        expected = {}  # (date, p, name) -> status in DB
        class_days = sorted(rnd.sample(range(1, days + 1), min(days, 22)))
        sid = dict(conn.execute("SELECT name,id FROM students"))
        jd = dict(conn.execute("SELECT name,joined_date FROM students"))
        for day in class_days:
            ds = f"{ym}-{day:02d}"
            for p in range(5):
                mode = rnd.random()
                if mode < 0.2: continue  # 整节无记录
                for n in base:
                    if jd[n] and ds < jd[n] and rnd.random() < 0.9: continue
                    if rnd.random() < 0.3: continue  # 部分学生无记录
                    st = rnd.choice(STATUSES)
                    app.save_record(ds, p, n, st); expected[(ds, p, n)] = st
        if delete_readd:
            victim = base[3]
            app.update_db_roster([x for x in base if x != victim])
            app.update_db_roster(base)
            # v1.4.0：移出是软删除，同名加回来记录恢复 → 期望值不变
            self.assertEqual(conn.execute("SELECT active FROM students WHERE name=?", (victim,)).fetchone()[0], 1)
            jd = dict(conn.execute("SELECT name,joined_date FROM students"))
        app.load_students_order()
        path, err = app.export_db_to_excel(ym, "24示例3班", today="2026-12-31")
        self.assertIsNone(err)
        cells, blocks = own_parse(path)
        # 每块：表头完整、人员 = 当前名单且各出现一次
        roster = [r[0] for r in conn.execute("SELECT name FROM students WHERE active=1 ORDER BY sort_order, id")]
        for d, ns, hdr in blocks:
            self.assertEqual(ns, roster); self.assertEqual(hdr, app.PERIOD_LABELS)
        # 逐格比对
        by_dp = {}
        for (ds, p, n) in expected: by_dp.setdefault((ds, p), set()).add(n)
        exported_dates = set()
        for (dl, p, n), v in cells.items():
            if not str(dl)[:2].isdigit() or "月" not in str(dl): raise AssertionError(("bad date header", dl, p, n))
            m, dd = int(dl[:2]), int(dl[3:5]); ds = f"{ym[:4]}-{m:02d}-{dd:02d}"; exported_dates.add(ds)
            if (ds, p, n) in expected:
                self.assertEqual(v, expected[(ds, p, n)], (ds, p, n))
            elif jd.get(n) and ds < jd[n]:
                self.assertIsNone(v, ("入班前应留空", ds, p, n))
            else:
                # 点过名的节里没记录的人、整节没点名的节：都是空课（与点名页一致）
                self.assertEqual(v, "空课", (ds, p, n))
        # 数据库里每条记录都出现在导出中（除全天空课的日期外）
        for (ds, p, n), st in expected.items():
            day_active = any(v != "空课" for (d2, _, _), v in expected.items() if d2 == ds)
            if day_active:
                self.assertIn(ds, exported_dates)
        if DOWNSTREAM:
            days, recs, vals = DOWNSTREAM(path)
            self.assertTrue(all(d[4] == len(roster) for d in days), "下游解析人数不符")
            want = sorted((n, st) for (ds, p, n), st in expected.items() if st in ("旷课", "病假", "事假"))
            got = sorted((r[2], r[3]) for r in recs)
            self.assertEqual(got, want)
        return len(roster), len(blocks)

    def test_45_plus_5_midmonth(self):  print(self.run_scenario(45, 5, 10))
    def test_45_plus_6(self):          print(self.run_scenario(45, 6, 15, seed=2))
    def test_60(self):                 print(self.run_scenario(55, 5, 8, seed=3))
    def test_100_reorder(self):        print(self.run_scenario(95, 5, 20, seed=4, reorder=True))
    def test_150_delete_readd(self):   print(self.run_scenario(140, 10, 12, seed=5, delete_readd=True))
    def test_october_31days(self):     print(self.run_scenario(48, 3, 31, ym="2026-10", seed=6))

    def test_missing_student_no_crash(self):
        self.fresh_db(); app.update_db_roster(["张三丰"])
        app.save_record("2026-09-01", 0, "不存在", "旷课")
        app.bulk_save_records("2026-09-01", 0, {"不存在": "旷课", "张三丰": "旷课"})
        self.assertEqual(app.get_records("2026-09-01", 0), {"张三丰": "旷课"})

    def test_soft_delete_keeps_records(self):
        c = self.fresh_db(); app.update_db_roster(["张三丰", "李四光"])
        app.save_record("2026-09-01", 0, "张三丰", "旷课"); app.update_db_roster(["李四光"])
        self.assertEqual(c.execute("SELECT COUNT(*) FROM records").fetchone()[0], 1)
        self.assertEqual(app.STUDENTS_ORDER, ["李四光"])

    def test_purge_deletes_records(self):
        c = self.fresh_db(); app.update_db_roster(["张三丰", "李四光"])
        app.save_record("2026-09-01", 0, "张三丰", "旷课")
        cid = app.get_current_class_id()
        sid = c.execute("SELECT id FROM students WHERE name='李四光'").fetchone()[0]
        app.apply_roster_changes(cid, [{"id": sid, "name": "李四光"}], purge_removed=True)
        self.assertEqual(c.execute("SELECT COUNT(*) FROM records").fetchone()[0], 0)
        self.assertEqual(c.execute("SELECT COUNT(*) FROM students").fetchone()[0], 1)

    def test_unmarked_is_noclass_and_left_is_blank(self):
        c = self.fresh_db(); app.update_db_roster(["张三丰", "李四光", "王五六"])
        app.save_record("2026-09-01", 0, "张三丰", "旷课")          # 只标了一个人
        app.save_record("2026-09-01", 1, "王五六", "出勤")
        c.execute("UPDATE students SET active=0, left_date='2026-09-02' WHERE name='王五六'"); c.commit()
        app.save_record("2026-09-03", 0, "张三丰", "病假")
        app.load_students_order()
        path, err = app.export_db_to_excel("2026-09", "测试班", today="2026-12-31")
        self.assertIsNone(err)
        cells, blocks = own_parse(path)
        self.assertEqual(cells[("09月01日", 0, "李四光")], "空课")   # 点过名但没记录 → 空课
        self.assertEqual(cells[("09月01日", 1, "王五六")], "出勤")   # 在班期间照常
        self.assertIsNone(cells[("09月03日", 0, "王五六")])          # 移出之后留空
        self.assertEqual(blocks[0][1], ["张三丰", "李四光", "王五六"])

    def test_future_dates_not_exported(self):
        self.fresh_db(); app.update_db_roster(["张三丰"])
        app.save_record("2026-10-20", 0, "张三丰", "旷课")
        app.save_record("2026-10-30", 0, "张三丰", "旷课")
        path, err = app.export_db_to_excel("2026-10", "测试班", today="2026-10-27")
        cells, _ = own_parse(path)
        self.assertTrue(any(k[0] == "10月20日" for k in cells))
        self.assertFalse(any(k[0] == "10月30日" for k in cells))

    # ---- 已移出同学的导出名单：移出日期晚于当月 1 号，或当月在本班有记录（v1.4.0 修复 1 号移出漏行） ----
    def _left_scenario(self, left_date, record_date=None, status="旷课"):
        c = self.fresh_db(); app.update_db_roster(["张三丰", "李四光", "王五六"])
        app.save_record("2026-10-01", 0, "张三丰", "出勤")
        app.save_record("2026-10-02", 0, "张三丰", "病假")
        if record_date:
            app.save_record(record_date, 0, "王五六", status)
        c.execute("UPDATE students SET active=0, left_date=? WHERE name='王五六'", (left_date,)); c.commit()
        app.load_students_order()
        path, err = app.export_db_to_excel("2026-10", "测试班", today="2026-12-31")
        self.assertIsNone(err)
        return own_parse(path)

    def test_left_on_first_with_record_is_exported(self):
        cells, blocks = self._left_scenario("2026-10-01", "2026-10-01", "旷课")
        for _, ns, _ in blocks:
            self.assertEqual(ns, ["张三丰", "李四光", "王五六"])     # 在班在前，移出在后
        self.assertEqual(cells[("10月01日", 0, "王五六")], "旷课")   # 当天的记录照填
        self.assertIsNone(cells[("10月01日", 1, "王五六")])          # 移出当天没记录的节留空（规则不变）
        self.assertIsNone(cells[("10月02日", 0, "王五六")])          # 移出之后留空

    def test_left_on_first_without_record_not_exported(self):
        cells, blocks = self._left_scenario("2026-10-01")
        for _, ns, _ in blocks:
            self.assertEqual(ns, ["张三丰", "李四光"])
        self.assertFalse(any(k[2] == "王五六" for k in cells))

    def test_left_midmonth_unchanged(self):
        cells, blocks = self._left_scenario("2026-10-02")
        for _, ns, _ in blocks:
            self.assertEqual(ns, ["张三丰", "李四光", "王五六"])
        self.assertEqual(cells[("10月01日", 0, "王五六")], "空课")   # 在班期间没记录 = 空课
        self.assertIsNone(cells[("10月02日", 0, "王五六")])

    def test_left_without_date_but_with_record_is_exported(self):
        cells, blocks = self._left_scenario(None, "2026-10-01", "事假")
        self.assertEqual(blocks[0][1], ["张三丰", "李四光", "王五六"])
        self.assertEqual(cells[("10月01日", 0, "王五六")], "事假")

if __name__ == "__main__":
    unittest.main()
