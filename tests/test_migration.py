# -*- coding: utf-8 -*-
"""覆盖安装升级：新装为空；v1.3.1 / v1.2.0 旧库升级后每条记录都在；失败回滚；备份失败不升级。"""
import os, sys, re, shutil, sqlite3, tempfile, unittest, glob
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tests", "fixtures"))
import main as app
import make_v131_db

REAL_V120 = os.environ.get("HMUDQ_REAL_V120_DB", "")  # 可选：本地真实 v1.2.0 库路径（不随仓库提供）


def raw_snapshot(path):
    """不经过 App，直接读库：学生集合 + 每条记录 (姓名, 日期, 节, 状态)。"""
    c = sqlite3.connect(path)
    studs = sorted(c.execute("SELECT id, name, COALESCE(sort_order,0) FROM students"))
    recs = sorted(c.execute("SELECT s.name, r.date, r.period, r.status FROM records r JOIN students s ON s.id=r.student_id"))
    ver = c.execute("PRAGMA user_version").fetchone()[0]
    tables = {r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    c.close()
    return studs, recs, ver, tables


class Base(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()
        app.close_db_connection() if app._db_conn else None
        app._db_conn = None
        app.DB_PATH = os.path.join(self.d, "attendance.db")
        app._MIGRATION_FAULT = None

    def tearDown(self):
        app._MIGRATION_FAULT = None
        if app._db_conn:
            app.close_db_connection()
        app._db_conn = None
        shutil.rmtree(self.d, ignore_errors=True)

    def reopen(self):
        app.close_db_connection(); app._db_conn = None


class FreshInstall(Base):
    def test_fresh_is_empty(self):
        info = app.init_db(); app.load_students_order()
        self.assertTrue(info["fresh"]); self.assertFalse(info["migrated"])
        self.assertEqual(app.STUDENTS_ORDER, [])
        self.assertEqual([c[1] for c in app.list_classes()], ["我的班级"])
        self.assertEqual(app.get_db_connection().execute("PRAGMA user_version").fetchone()[0], 2)
        self.assertEqual(app.get_db_connection().execute("SELECT COUNT(*) FROM records").fetchone()[0], 0)
        self.assertEqual(app.list_backups(), [])          # 空库不做升级备份


class UpgradeV131(Base):
    def make_old(self):
        _, info = make_v131_db.make(app.DB_PATH)
        return info

    def test_upgrade_keeps_every_record(self):
        gen = self.make_old()
        s0, r0, v0, t0 = raw_snapshot(app.DB_PATH)
        self.assertEqual(v0, 0); self.assertNotIn("classes", t0)
        self.assertGreater(len(r0), 1800)
        info = app.init_db(); app.load_students_order()
        self.assertTrue(info["migrated"])
        self.assertTrue(info["backup"] and os.path.exists(info["backup"]))
        self.reopen()
        s1, r1, v1, t1 = raw_snapshot(app.DB_PATH)
        self.assertEqual(r1, r0, "升级后考勤记录必须逐条一致")
        self.assertEqual(s1, s0, "学生 id / 名字 / 排序不变")
        self.assertEqual(v1, 2); self.assertIn("classes", t1)
        app.init_db(); app.load_students_order()
        self.assertEqual(app.get_class_name(), "临床3班")
        self.assertEqual(len(app.STUDENTS_ORDER), gen["students"])
        # 备份内容 = 升级前的库
        sb, rb, vb, _ = raw_snapshot(info["backup"])
        self.assertEqual((sb, rb, vb), (s0, r0, 0))
        # 部分标记的课节数：生成器知道有几节只标了缺勤
        self.assertEqual(info["partial_slots"], gen["partial"])
        self.__class__.last = (len(r0), gen)

    def test_second_start_is_noop(self):
        self.make_old()
        app.init_db(); self.reopen()
        n_backups = len(app.list_backups())
        before = raw_snapshot(app.DB_PATH)
        info2 = app.init_db()
        self.assertFalse(info2["migrated"]); self.assertFalse(info2["fresh"])
        self.reopen()
        self.assertEqual(raw_snapshot(app.DB_PATH), before)
        self.assertEqual(len(app.list_backups()), n_backups)

    def test_upgrade_page_and_export_agree(self):
        """升级后：点名页显示 == 导出 Excel（没记录 → 空课）。"""
        self.make_old()
        app.init_db(); app.load_students_order()
        import openpyxl
        path, err = app.export_db_to_excel("2026-10", app.get_class_name(), today="2026-10-27")
        self.assertIsNone(err)
        self.assertEqual(os.path.basename(path), "10月临床3班上课出勤.xlsx")
        ws = openpyxl.load_workbook(path).active
        checked = 0
        for r in range(1, ws.max_row + 1):
            for c0 in (1, 8, 15):
                if ws.cell(r, c0).value != "姓名":
                    continue
                dl = ws.cell(r, c0 + 1).value
                ds = f"2026-{int(dl[:2]):02d}-{int(dl[3:5]):02d}"
                rr = r + 1
                while rr <= ws.max_row and ws.cell(rr, c0).value not in (None, "姓名"):
                    n = ws.cell(rr, c0).value
                    for p in range(5):
                        page = app.effective_statuses([(n, app.STUDENT_JOINED.get(n))], app.get_records(ds, p), ds)[n]
                        v = ws.cell(rr, c0 + 2 + p).value
                        self.assertEqual(v if v is not None else app.NOT_JOINED, page, (ds, p, n))
                        checked += 1
                    rr += 1
        self.assertGreater(checked, 1000)


    def test_upgrade_never_fills_partial_slots(self):
        """升级不补记、不弹窗：考勤条数与升级前完全一致，部分标记课节里没记录的人是空课。"""
        info = self.make_old()
        _, before, _, _ = raw_snapshot(app.DB_PATH)
        res = app.init_db(); app.load_students_order()
        self.assertTrue(res["migrated"]); self.assertGreater(res["partial_slots"], 0)
        _, after, _, _ = raw_snapshot(app.DB_PATH)
        self.assertEqual(sorted(after), sorted(before))
        self.assertEqual(len(after), info["records"])
        self.assertFalse(hasattr(app, "fill_partial_slots"))
        cname, d, p, n = app.partial_slots(until="2026-10-27")[0]
        roster = [(x, app.STUDENT_JOINED.get(x)) for x in app.STUDENTS_ORDER]
        eff = app.effective_statuses(roster, app.get_records(d, p), d)
        self.assertGreaterEqual(list(eff.values()).count("空课"), n)

    def test_no_partial_dialog_in_ui(self):
        src = open(os.path.join(ROOT, "main.py"), encoding="utf-8").read()
        for s in ("open_partial_slots_dialog", "补记为出勤", "fill_partial_slots"):
            self.assertNotIn(s, src)


class UpgradeV120Real(Base):
    @unittest.skipUnless(REAL_V120 and os.path.exists(REAL_V120), "没有真实 v1.2.0 库")
    def test_real_v120_db(self):
        shutil.copy(REAL_V120, app.DB_PATH)
        s0, r0, _, _ = raw_snapshot(app.DB_PATH)
        info = app.init_db(); self.reopen()
        s1, r1, v1, _ = raw_snapshot(app.DB_PATH)
        self.assertTrue(info["migrated"])
        self.assertEqual(r1, r0); self.assertEqual(len(r1), len(r0))
        self.assertEqual([x[1] for x in s1], [x[1] for x in s0]); self.assertEqual(v1, 2)
        print(f"\n[v1.2.0 实库] 学生 {len(s1)}，记录 {len(r1)}，部分标记课节 {info['partial_slots']}")


class Rollback(Base):
    def _fault(self, where):
        make_v131_db.make(app.DB_PATH)
        before = raw_snapshot(app.DB_PATH)
        app._MIGRATION_FAULT = where
        with self.assertRaises(app.MigrationError) as cm:
            app.init_db()
        self.assertTrue(cm.exception.backup_path and os.path.exists(cm.exception.backup_path))
        self.reopen()
        self.assertEqual(raw_snapshot(app.DB_PATH), before, "失败后数据库必须原样")
        # 去掉故障后再启动，能正常升级
        app._MIGRATION_FAULT = None
        info = app.init_db(); self.reopen()
        self.assertTrue(info["migrated"])
        self.assertEqual(raw_snapshot(app.DB_PATH)[1], before[1])

    def test_fault_after_copy(self):    self._fault("after_copy")
    def test_fault_verify_fails(self):  self._fault("before_verify")

    def test_backup_failure_stops_upgrade(self):
        make_v131_db.make(app.DB_PATH)
        before = raw_snapshot(app.DB_PATH)
        orig = app.make_backup
        def boom(*a, **k): raise OSError("磁盘已满")
        app.make_backup = boom
        try:
            with self.assertRaises(app.MigrationError):
                app.init_db()
        finally:
            app.make_backup = orig
        self.reopen()
        self.assertEqual(raw_snapshot(app.DB_PATH), before)

    def test_records_without_students_table(self):
        c = sqlite3.connect(app.DB_PATH)
        c.execute("CREATE TABLE records (id INTEGER PRIMARY KEY, student_id INT, date TEXT, period INT, status TEXT)")
        c.execute("INSERT INTO records VALUES (1, 1, '2026-10-01', 0, '出勤')"); c.commit(); c.close()
        with self.assertRaises(app.MigrationError):
            app.init_db()
        self.reopen()
        c = sqlite3.connect(app.DB_PATH)
        self.assertEqual(c.execute("SELECT COUNT(*) FROM records").fetchone()[0], 1); c.close()


class SourceRules(unittest.TestCase):
    def test_never_deletes_db(self):
        src = open(os.path.join(ROOT, "main.py"), encoding="utf-8").read()
        self.assertIsNone(re.search(r"os\.(remove|unlink)\(\s*DB_PATH", src))
        self.assertIsNone(re.search(r"shutil\.rmtree", src))

    def test_db_path_function_unchanged(self):
        """get_db_path 必须与 v1.3.1 一字不差（覆盖安装找到同一个库）。"""
        def f(s):
            i = s.index("def get_db_path"); return s[i:s.index("\n\n", i)]
        new = open(os.path.join(ROOT, "main.py"), encoding="utf-8").read()
        old = open(os.path.join(ROOT, "tests", "fixtures", "main_v131.py"), encoding="utf-8").read()
        self.assertEqual(f(new), f(old))


if __name__ == "__main__":
    unittest.main()
